"""Audit retained YNP1 source traces without assigning unverified strain aliases."""
from pathlib import Path
from datetime import datetime, timezone
import hashlib
import json
from openpyxl import load_workbook

P = Path(__file__).resolve().parent
ROOT = P.parents[2]
book = P / '41589_2026_2282_MOESM13_ESM.xlsx'
inventory = P / 'massive-inventory.json'
files = json.loads(inventory.read_text(encoding='utf-8'))['files']
w = load_workbook(book, read_only=True, data_only=True)
traces = []
for sheet in w:
    rows = list(sheet.iter_rows(values_only=True))
    for col, acquisition in enumerate(rows[1]):
        if not isinstance(acquisition, str) or not acquisition.endswith(('.mzXML', '.mzML')):
            continue
        values = []
        for r in rows[3:]:
            try:
                values.append((float(r[col]), float(r[col + 1])))
            except (ValueError, TypeError, IndexError):
                continue
        peak = max(values, key=lambda x: x[1]) if values else None
        stem = Path(acquisition).stem
        matches = [f for f in files if Path(f['path']).stem == stem]
        # These literal inserted labels were observed in the retained FTP inventory.
        # Preserve the original names and call these candidates, not conversion proof.
        insertions = {'SC YNP1':'SC_YNP1_', 'EV 3F YNP1':'EX_3F_YNP1_',
                      'EV 4F YNP1':'EX_4F_YNP1_', 'KP YNP1':'KP_YNP1_'}
        inserted_label = insertions.get(rows[0][col])
        candidates = []
        if inserted_label:
            for f in files:
                raw_stem = Path(f['path']).stem
                if inserted_label in raw_stem and raw_stem.replace(inserted_label, '', 1) == stem:
                    candidates.append(f)
        traces.append(dict(sheet=sheet.title, source_cell=f'{sheet.cell(2,col+1).column_letter}2',
                           source_label=rows[0][col], acquisition=acquisition,
                           numeric_trace_rows=len(values), maximum_rt_intensity=peak,
                           exact_stem_inventory_matches=matches,
                           observed_raw_inserted_label=inserted_label,
                           label_removal_locator_candidates=candidates))
w.close()
inputs = [book, inventory, P/'supplement.txt', P/'gadar-lopez-thesis.txt',
          P/'ynp1-source-anchor-audit.json']
result = dict(run_id='ynp1-retained-measurement-crosswalk-2026-10-09',
              created_utc=datetime.now(timezone.utc).isoformat(), state='completed',
              command='python data/interim/black-yeast-screen-2026-10-09/check_ynp1_measurement_crosswalk.py',
              inputs=[dict(path=x.relative_to(ROOT).as_posix(),sha256=hashlib.sha256(x.read_bytes()).hexdigest()) for x in inputs],
              traces=traces,
              interpretation='Workbook source labels establish panel/acquisition host-condition correspondence. Exact acquisition stems locate retained deposit entries, not native conversion or spectrum joins. C661/C664/785 to sYNP19221/19222/19223 remains a source-context inference, not an explicit identifier crosswalk. Peak maxima are trace summaries, not independent chemical validation. No producer sequence comparison or admission performed.')
out=P/'ynp1-measurement-crosswalk.json'
out.write_text(json.dumps(result,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'state':'completed','traces':len(traces),'traces_with_inventory_matches':sum(bool(x['exact_stem_inventory_matches']) for x in traces),'traces_with_label_removal_candidates':sum(bool(x['label_removal_locator_candidates']) for x in traces),'output':out.relative_to(ROOT).as_posix()}))
